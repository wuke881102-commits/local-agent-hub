"""SlidePlan → .pptx。**这个模块不打模型、不联网。**

边界这么划是为了它能脱离模型单测：给一份 SlidePlan，断言出来的 .pptx 有几页、
每页几个形状、文字落在哪个形状里。识别那半边（会打模型）在 slide_rebuild.py。
"""
from __future__ import annotations

import logging
import math
import shutil
from dataclasses import dataclass, field
from pathlib import Path

log = logging.getLogger("services.pptx_render")

RGB = tuple[int, int, int]

# ── 配色：Lumen-light 设计系统（浅色版）───────────────────────────────
# 与 HTML 产物同一套令牌（app/html/design_system.py），两种格式拿出去是一家的。
# 按图还原也用这套：还原的是**版面结构**，颜色统一走品牌色，原图页紧跟在后面核对。
BRAND_50: RGB = (0xE3, 0xF5, 0xEA)
BRAND_200: RGB = (0x90, 0xD4, 0xA8)
BRAND_400: RGB = (0x30, 0xAA, 0x60)
BRAND_500: RGB = (0x00, 0xAA, 0x4F)
BRAND_600: RGB = (0x00, 0x8E, 0x43)
BRAND_700: RGB = (0x00, 0x68, 0x45)
ACCENT_LIME: RGB = (0x82, 0xBC, 0x00)
SURFACE_PAGE: RGB = (0xF9, 0xFA, 0xFC)
SURFACE_ELEVATED: RGB = (0xFF, 0xFF, 0xFF)
SURFACE_SUBTLE: RGB = (0xF0, 0xF1, 0xF3)
BORDER_DEFAULT: RGB = (0xDD, 0xE3, 0xEA)
BORDER_SUBTLE: RGB = (0xE8, 0xEC, 0xF0)
TEXT_PRIMARY: RGB = (0x42, 0x46, 0x4A)
TEXT_SECONDARY: RGB = (0x55, 0x5B, 0x61)
TEXT_TERTIARY: RGB = (0x73, 0x7A, 0x82)
WHITE: RGB = (0xFF, 0xFF, 0xFF)

# 字体：规范的 sans 栈在 Windows 上第一个落地的是 Segoe UI；中文字形规范没写，
# 用系统自带的微软雅黑（Inter 没有中文，写了也会被替换成宋体）。
FONT_SANS = "Segoe UI"
FONT_MONO = "Consolas"
FONT_EA = "Microsoft YaHei"

# 圆角（英寸）。规范是 px，按 16:9 页面 ≈ 1280px 宽换算：1px ≈ 1/96 in。
RADIUS_LG = 8 / 96
RADIUS_XL = 12 / 96
RADIUS_2XL = 16 / 96


@dataclass
class Block:
    kind: str = "box"                 # box | group | note | callout
    label: str = ""
    sub: str = ""
    items: list[str] = field(default_factory=list)
    value: str = ""                   # callout 专用，如 "↑80%"
    accent: str = "plain"             # blue | green | plain


@dataclass
class Column:
    index: str = ""                   # "00" / "01"，没有编号就空
    heading: str = ""
    blocks: list[Block] = field(default_factory=list)
    # 这一栏在**原照片**里的归一化位置 [x0,y0,x1,y1]，由视觉模型给出。
    # 渲染用不到它，切块用得到 —— 放在这里是因为它和栏一一对应。
    box: tuple[float, float, float, float] | None = None


@dataclass
class Band:
    role: str = "columns"             # flow（横向流）| columns（分栏）
    columns: list[Column] = field(default_factory=list)


@dataclass
class SlidePlan:
    is_slide: bool = True
    title: str = ""
    subtitle: str = ""
    theme: str = "dark"
    bands: list[Band] = field(default_factory=list)
    notes: str = ""                   # 全文转写，进演讲者备注


# ── 讲稿模式（文字型文档）──────────────────────────────────────────

TALK_KINDS = ("bullets", "metrics", "table", "closing")


@dataclass
class TalkSlide:
    """讲稿的一页。和 SlidePlan 分开：那边是「还原一张展板」，这边是「讲一个主题」，
    版面逻辑完全不同，硬塞进同一个结构只会让两边都别扭。"""
    kind: str = "bullets"              # bullets | metrics | table | closing | appendix
    title: str = ""
    bullets: list[str] = field(default_factory=list)
    metrics: list[dict] = field(default_factory=list)      # [{"value": "80%", "label": "…"}]
    columns: list[str] = field(default_factory=list)
    rows: list[list[str]] = field(default_factory=list)
    image: str | None = None           # 图片编号，如 "I3"；渲染时换成本地文件
    notes: str = ""                    # 原文细节，进演讲者备注
    sources: list[str] = field(default_factory=list)       # 引用了哪几节，如 ["S2","S5"]


@dataclass
class Talk:
    title: str = ""
    subtitle: str = ""
    slides: list[TalkSlide] = field(default_factory=list)
    dropped: list[str] = field(default_factory=list)       # 校验时丢掉了什么、为什么


@dataclass
class Rect:
    """英寸坐标。python-pptx 用 EMU，转换在画图那一层做，这里保持可读。"""
    x: float
    y: float
    w: float
    h: float


def _weight(band: Band) -> int:
    """一条带的内容量。用块数而不是字数 —— 块数直接决定要堆多少个形状。"""
    return max(1, sum(len(c.blocks) for c in band.columns))


def band_heights(bands: list[Band], total: float, gap: float) -> list[float]:
    """按内容量把总高分给各条带。

    空带也给正高度：0 高的形状 python-pptx 会拒绝，而「这条带是空的」本身
    是识别结果的一部分，不该让整页画不出来。
    """
    if not bands:
        return []
    usable = total - gap * (len(bands) - 1)
    weights = [_weight(b) for b in bands]
    s = sum(weights)
    return [usable * w / s for w in weights]


MIN_BLOCK_H = 0.28          # 再矮就放不下 9pt 的字了


def stack_rects(area: Rect, n: int, gap: float, min_h: float,
                weights: list[float] | None = None) -> list[Rect]:
    """在 area 里竖向堆 n 个块，**保证全部落在 area 之内**。

    放不下就少堆几个（调用方把剩下的写进备注），而不是硬塞。
    上一版这里是 `max(0.3, 可用高/n)` —— 那个下限治了「0 高形状 python-pptx
    会拒绝」，却造出了更糟的病：块一多，堆叠总高超过本带高度，直接压到下一条带
    身上，两带文字叠在一起。真文档跑出来才看见，所有单测都是绿的。
    """
    if n <= 0 or area.h <= 0:
        return []
    fit = n
    while fit > 1 and (area.h - gap * (fit - 1)) / fit < min_h:
        fit -= 1
    usable = area.h - gap * (fit - 1)
    w = [max(1e-6, x) for x in (weights or [1.0] * n)[:fit]]
    spare = usable - min_h * fit
    if spare <= 0:
        # 连保底都给不起，只能均分
        h = max(0.01, usable / fit)
        heights = [h] * fit
    else:
        # 先每块给 min_h 保底，剩下的按权重分 —— 字多的框拿得多。
        # 8.3 一律均分，一栏里有个长句时它和「高置信风险」一样高，只能溢出。
        total = sum(w)
        heights = [min_h + spare * x / total for x in w]
    out, y = [], area.y
    for h in heights:
        out.append(Rect(area.x, y, area.w, h))
        y += h + gap
    return out


def column_rects(band: Band, area: Rect, gap: float) -> list[Rect]:
    """把一条带的横向空间均分给各栏。"""
    n = len(band.columns)
    if n == 0:
        return []
    w = (area.w - gap * (n - 1)) / n
    return [Rect(area.x + i * (w + gap), area.y, w, area.h) for i in range(n)]


MIN_FONT_PT = 9
MAX_FONT_PT = 18

# python-pptx 自选图形文本框的默认内边距：左右各 0.1"，上下各 0.05"。
TEXT_MARGIN_X = 0.2
TEXT_MARGIN_Y = 0.1
LINE_SPACING = 1.2          # 行高 ≈ 字号 × 1.2
LATIN_EM = 0.55             # 拉丁字母约半个中文字宽


def _text_em(text: str) -> float:
    """一段文字的宽度，以 em（=字号）为单位。中日韩字符按 1em，其余按 0.55em。"""
    return sum(1.0 if ord(c) >= 0x2E80 else LATIN_EM for c in text)


def text_lines(text: str, width: float, pt: int) -> int:
    """这段文字放进宽 ``width`` 英寸的框、用 ``pt`` 号字，会折成几行。

    换行符分段，每段至少一行。这是 fit_font_size 和 stack_rects 共用的那把尺子。
    """
    per_line = max(1.0, (max(0.05, width - TEXT_MARGIN_X) * 72) / pt)
    return sum(max(1, math.ceil(_text_em(para) / per_line))
               for para in (text or "").split("\n"))


def fit_font_size(text: str, width: float, height: float, max_pt: int = MAX_FONT_PT) -> int:
    """能让这段文字装进框的最大字号（pt），下限 9pt。

    8.3 这里是「面积 ÷ 字数 开根号」，没算形状内边距、行距和换行，一律估大了。
    PowerPoint 对装不下的文字是**上下对称**溢出的，往上就压到了栏标题 ——
    装机版上肉眼看到的，单测全绿。现在逐档试字号，按真实折行后的高度判断。

    到了下限仍装不下就返回下限：低于 9pt 投影上没人看得见，那时该做的是给框
    更多高度（见 stack_rects 的权重），而不是继续缩字。
    """
    usable = max(0.05, height - TEXT_MARGIN_Y) * 72
    for pt in range(max_pt, MIN_FONT_PT - 1, -1):
        if text_lines(text, width, pt) * pt * LINE_SPACING <= usable:
            return pt
    return MIN_FONT_PT


def block_need(text: str, width: float, pt: int = 12) -> float:
    """按参考字号，这段文字需要多高的框（英寸）。给 stack_rects 当权重用。"""
    return text_lines(text, width, pt) * pt * LINE_SPACING / 72 + TEXT_MARGIN_Y


# ── 画图 ──────────────────────────────────────────────────────────
from pptx import Presentation                       # noqa: E402
from pptx.dml.color import RGBColor                 # noqa: E402
from pptx.enum.shapes import MSO_SHAPE              # noqa: E402
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN                 # noqa: E402
from pptx.oxml.ns import qn                         # noqa: E402
from pptx.util import Emu, Inches, Pt               # noqa: E402
from lxml import etree                              # noqa: E402

SLIDE_W = 13.333          # 16:9
SLIDE_H = 7.5
MARGIN = 0.45
TITLE_H = 1.05
GAP = 0.16
BLANK_LAYOUT = 6          # 默认模板里 6 号是全空白版式
PHOTO_MAX_PX = 1920       # 嵌入原图页的长边上限，见 _embed_ready


@dataclass
class DeckItem:
    plan: SlidePlan
    photo: Path | None = None


def _blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[BLANK_LAYOUT])


def _style_run(run, *, size: float, color: RGB, bold: bool = False, mono: bool = False) -> None:
    """统一的字形：拉丁用 Segoe UI（或 Consolas），中文用微软雅黑。

    python-pptx 的 font.name 只写 a:latin；不补 a:ea，中文会落到主题的东亚字体
    （默认模板是宋体），一页里中英文两种风格。
    """
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = RGBColor(*color)
    f.name = FONT_MONO if mono else FONT_SANS
    rpr = run._r.get_or_add_rPr()
    # 标成中文，PowerPoint 才按中文的避头尾规则断行 —— 否则「、」「）」会落到行首。
    rpr.set("lang", "zh-CN")
    rpr.set("altLang", "en-US")
    ppr = run._r.getparent().get_or_add_pPr()
    ppr.set("eaLnBrk", "1")
    ppr.set("hangingPunct", "1")
    for old in rpr.findall(qn("a:ea")):
        rpr.remove(old)
    ea = etree.Element(qn("a:ea"))
    ea.set("typeface", FONT_EA)
    rpr.find(qn("a:latin")).addnext(ea)


def _text(slide, rect: Rect, text: str, *, size: int, color: RGB,
          bold: bool = False, align=PP_ALIGN.LEFT, mono: bool = False):
    box = slide.shapes.add_textbox(Inches(rect.x), Inches(rect.y),
                                   Inches(rect.w), Inches(rect.h))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    _style_run(run, size=size, color=color, bold=bold, mono=mono)
    return box


def _flat(shape) -> None:
    """去掉主题默认的阴影。规范的 shadow-sm 只是一圈极淡的描边感，用细边框表达。"""
    shape.shadow.inherit = False


def _card(slide, rect: Rect, *, fill: RGB = SURFACE_ELEVATED, border: RGB | None = BORDER_SUBTLE,
          radius: float = RADIUS_2XL):
    """规范里的卡片：白底 + 圆角 + 细边框。"""
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                   Inches(rect.x), Inches(rect.y), Inches(rect.w), Inches(rect.h))
    shape.adjustments[0] = min(0.5, radius / max(0.01, min(rect.w, rect.h)))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(*fill)
    if border is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = RGBColor(*border)
        shape.line.width = Pt(0.75)
    _flat(shape)
    return shape


def _gradient(shape, stops: list[tuple[float, RGB]], angle: float) -> None:
    """线性渐变。angle 按 CSS 的习惯：180 = 从上到下，135 = 左上到右下。

    python-pptx 的渐变只给两个色标、角度方向也和 CSS 不同，直接写 a:gradFill。
    """
    shape.fill.gradient()
    grad = shape._element.spPr.find(qn("a:gradFill"))
    gs_lst = grad.find(qn("a:gsLst"))
    for gs in list(gs_lst):
        gs_lst.remove(gs)
    for pos, rgb in stops:
        gs = etree.SubElement(gs_lst, qn("a:gs"))
        gs.set("pos", str(int(pos * 100000)))
        clr = etree.SubElement(gs, qn("a:srgbClr"))
        clr.set("val", "%02X%02X%02X" % rgb)
    lin = grad.find(qn("a:lin"))
    if lin is None:
        lin = etree.SubElement(grad, qn("a:lin"))
    # CSS 0deg 朝上、顺时针；DrawingML 0 朝右、顺时针，单位 1/60000 度。
    lin.set("ang", str(int(((angle - 90) % 360) * 60000)))
    lin.set("scaled", "0")


def _alpha(shape, rgb: RGB, opacity: float) -> None:
    """半透明纯色填充（封面上的圆形装饰）。"""
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(*rgb)
    clr = shape._element.spPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
    a = etree.SubElement(clr, qn("a:alpha"))
    a.set("val", str(int(opacity * 100000)))


def _page(slide) -> None:
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(*SURFACE_PAGE)


def _title_bar(slide, x: float, y: float, h: float) -> None:
    """标题前的品牌色竖条：和 HTML 卡片标题的 h2::before 同一个做法（brand-400→600 渐变）。"""
    bar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y),
                                 Inches(0.07), Inches(h))
    bar.adjustments[0] = 0.5
    _gradient(bar, [(0, BRAND_400), (1, BRAND_600)], 180)
    bar.line.fill.background()
    _flat(bar)


def _block_text(b: Block) -> str:
    """一个块画出来是什么字 —— 和 _draw_block 的取法保持一致，权重才估得准。"""
    body = b.value or b.label
    sub = b.sub or ("、".join(b.items) if b.items else "")
    if not sub and b.value and b.label:
        sub = b.label
    return f"{body}\n{sub}" if sub else body


def _draw_block(slide, b: Block, rect: Rect) -> None:
    if b.kind == "note":
        _text(slide, rect, b.label,
              size=fit_font_size(b.label, rect.w, rect.h), color=TEXT_SECONDARY)
        return

    # 原图里强调过的块（蓝/绿底）→ 品牌绿实底白字；普通块 → 白卡片深字。
    strong = b.accent != "plain"
    shape = _card(slide, rect, fill=BRAND_500 if strong else SURFACE_ELEVATED,
                  border=None if strong else BORDER_SUBTLE, radius=RADIUS_XL)
    main_color = WHITE if strong else (BRAND_600 if b.value else TEXT_PRIMARY)
    sub_color = BRAND_50 if strong else TEXT_TERTIARY

    body = b.value or b.label
    sub = b.sub or ("、".join(b.items) if b.items else "")
    if not sub and b.value and b.label:
        # callout：value 是那个大数字，label 是它的说明。value 占了主文字位之后，
        # label 就没有去处了 —— 不接这一句，「↑80%」画得出来而「客服响应效率」
        # 凭空消失，且所有断言都是绿的。
        sub = b.label
    size = fit_font_size(f"{body}\n{sub}" if sub else body, rect.w, rect.h)
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = body
    _style_run(r, size=size, color=main_color, bold=True)
    if sub:
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        r2 = p2.add_run()
        r2.text = sub
        _style_run(r2, size=max(MIN_FONT_PT, size - 2), color=sub_color)


def _draw_plan(slide, plan: SlidePlan) -> None:
    _page(slide)
    _title_bar(slide, MARGIN, 0.3, 0.46)
    _text(slide, Rect(MARGIN + 0.2, 0.25, SLIDE_W - 2 * MARGIN - 0.2, 0.55),
          plan.title, size=26, color=TEXT_PRIMARY, bold=True)
    if plan.subtitle:
        _text(slide, Rect(MARGIN + 0.2, 0.82, SLIDE_W - 2 * MARGIN - 0.2, 0.35),
              plan.subtitle, size=12, color=TEXT_TERTIARY)

    area_y = TITLE_H + 0.15
    area_h = SLIDE_H - area_y - MARGIN
    for band, h in zip(plan.bands, band_heights(plan.bands, area_h, GAP)):
        area = Rect(MARGIN, area_y, SLIDE_W - 2 * MARGIN, h)
        for ci, (col, crect) in enumerate(zip(band.columns, column_rects(band, area, GAP))):
            y = crect.y
            if col.index or col.heading:
                box = slide.shapes.add_textbox(Inches(crect.x), Inches(y), Inches(crect.w), Inches(0.3))
                tf = box.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                if col.index:
                    r = p.add_run()
                    r.text = col.index + " "
                    _style_run(r, size=12, color=BRAND_500, bold=True, mono=True)
                if col.heading:
                    r = p.add_run()
                    r.text = col.heading
                    _style_run(r, size=12, color=TEXT_PRIMARY, bold=True)
                y += 0.34
            inner = Rect(crect.x, y, crect.w, crect.y + crect.h - y)
            rects = stack_rects(inner, len(col.blocks), GAP, MIN_BLOCK_H,
                                weights=[block_need(_block_text(b), crect.w)
                                         for b in col.blocks])
            for b, r in zip(col.blocks, rects):
                _draw_block(slide, b, r)
            dropped = len(col.blocks) - len(rects)
            if dropped > 0:
                # 放不下的块不画，但要让人知道它们存在 —— 全文在备注里。
                log.debug("第 %d 栏有 %d 块放不下，已略去（全文在备注）", ci + 1, dropped)
        area_y += h + GAP


def _embed_ready(photo: Path, workdir: Path) -> Path:
    """必要时把照片缩到 PHOTO_MAX_PX 再嵌入，返回实际要嵌的文件。

    原图页是拿来核对识别结果的，屏幕上看清每个字就够，不需要手机的 4032px 原片。
    实测 10 张图的 deck 有 37 MB、22 张 80 MB —— 那种体积发不出去也打不开。
    本来就小的图原样用，不再压一道白掉画质。
    """
    from PIL import Image as _Image

    try:
        with _Image.open(photo) as im:
            if max(im.size) <= PHOTO_MAX_PX:
                return photo
            im = im.convert("RGB")
            im.thumbnail((PHOTO_MAX_PX, PHOTO_MAX_PX))
            workdir.mkdir(parents=True, exist_ok=True)
            out = workdir / f"{photo.stem}_embed.jpg"
            im.save(out, format="JPEG", quality=88)
            return out
    except Exception as e:  # noqa: BLE001
        log.debug("缩图失败，直接嵌原图：%s: %s", type(e).__name__, e)
        return photo


def _draw_photo(slide, photo: Path, caption: str, workdir: Path) -> None:
    """整页铺原图，等比缩放居中。原图页是用来核对识别结果的，必须看得清。"""
    from PIL import Image as _Image

    src = _embed_ready(photo, workdir)
    with _Image.open(src) as im:
        iw, ih = im.size
    scale = min((SLIDE_W - 0.6) / iw, (SLIDE_H - 1.0) / ih)
    w, h = iw * scale, ih * scale
    _page(slide)
    slide.shapes.add_picture(str(src), Inches((SLIDE_W - w) / 2), Inches(0.75),
                             Inches(w), Inches(h))
    _text(slide, Rect(MARGIN, 0.2, SLIDE_W - 2 * MARGIN, 0.4),
          caption, size=12, color=TEXT_TERTIARY)


def render_deck(items: list[DeckItem], out_path: Path) -> Path:
    """把若干 (还原计划, 原图) 渲染成一份 .pptx 并返回路径。"""
    prs = Presentation()
    prs.slide_width = Emu(int(Inches(SLIDE_W)))
    prs.slide_height = Emu(int(Inches(SLIDE_H)))

    import tempfile

    workdir = Path(tempfile.mkdtemp(prefix="lah-embed-"))
    for i, item in enumerate(items, 1):
        plan = item.plan
        if plan.is_slide:
            s = _blank(prs)
            _draw_plan(s, plan)
            if plan.notes:
                s.notes_slide.notes_text_frame.text = plan.notes
        if item.photo is not None:
            ps = _blank(prs)
            _draw_photo(ps, item.photo, f"原图 · 第 {i} 张　{plan.title}".strip(), workdir)
            if not plan.is_slide and plan.notes:
                ps.notes_slide.notes_text_frame.text = plan.notes

    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    shutil.rmtree(workdir, ignore_errors=True)
    return out_path


# ── 讲稿模式渲染 ──────────────────────────────────────────────────

TALK_TITLE_PT = 28
TALK_BODY_MAX_PT = 24         # 要点少时字可以大一些；多时由 fit_font_size 往下缩
TALK_TITLE_H = 0.9
IMAGE_SHARE = 0.40            # 带图的页，图占正文宽度的比例
CARD_PAD = 0.25               # 卡片内边距（规范 space-6 = 24px）
CELL_PAD_X = 16 / 96          # 表格单元格内边距（规范 space-3 space-4）
CELL_PAD_Y = 12 / 96


def _new_talk_slide(prs):
    s = _blank(prs)
    _page(s)
    return s


def _talk_title(slide, title: str) -> None:
    _title_bar(slide, MARGIN, 0.44, 0.5)
    _text(slide, Rect(MARGIN + 0.2, 0.3, SLIDE_W - 2 * MARGIN - 0.2, TALK_TITLE_H),
          title, size=TALK_TITLE_PT, color=TEXT_PRIMARY, bold=True)


def _talk_cover(prs, title: str, subtitle: str) -> None:
    """封面 = 规范里的头部 Banner：品牌绿 135° 渐变 + 两个半透明圆形装饰 + 白字。"""
    s = _blank(prs)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(SLIDE_W), Inches(SLIDE_H))
    _gradient(bg, [(0, BRAND_500), (0.6, BRAND_600), (1, BRAND_700)], 135)
    bg.line.fill.background()
    _flat(bg)
    for x, y, d, rgb, op in ((SLIDE_W - 3.6, -1.9, 5.2, WHITE, 0.10),
                             (-1.2, SLIDE_H - 2.6, 5.8, ACCENT_LIME, 0.18)):
        c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
        _alpha(c, rgb, op)
        c.line.fill.background()
        _flat(c)
    w = SLIDE_W - 2 * (MARGIN + 0.5)
    _text(s, Rect(MARGIN + 0.5, 2.5, w, 1.4), title or "讲稿",
          size=fit_font_size(title or "讲稿", w, 1.4, max_pt=44), color=WHITE, bold=True)
    if subtitle:
        _text(s, Rect(MARGIN + 0.5, 3.95, w, 0.8), subtitle,
              size=fit_font_size(subtitle, w, 0.8, max_pt=20), color=BRAND_50)


def _bullets_height(bullets: list[str], width: float, pt: int) -> float:
    """这组要点在 pt 字号下实际要多高（英寸），含段后距。"""
    lines = sum(text_lines(f"•  {b}", width, pt) for b in bullets)
    return (lines * pt * LINE_SPACING + len(bullets) * pt * 0.5) / 72 + TEXT_MARGIN_Y


def _bullet_box(slide, rect: Rect, bullets: list[str]) -> None:
    """一组要点放进一张白卡片，字号按真实折行算到装得下为止。圆点用品牌绿。

    卡片高度跟着内容走（和 HTML 的卡片一样）：三条要点配一张占满整页的白卡，
    下半截空着，看起来像没做完。
    """
    inner_w = rect.w - 2 * CARD_PAD
    lines = [f"•  {b}" for b in bullets]
    size = fit_font_size("\n".join(lines), inner_w, rect.h - 2 * CARD_PAD, max_pt=TALK_BODY_MAX_PT)
    need = _bullets_height(bullets, inner_w, size)
    h = min(rect.h, need + 2 * CARD_PAD)
    _card(slide, Rect(rect.x, rect.y, rect.w, h))
    inner = Rect(rect.x + CARD_PAD, rect.y + CARD_PAD, inner_w, h - 2 * CARD_PAD)
    box = slide.shapes.add_textbox(Inches(inner.x), Inches(inner.y), Inches(inner.w), Inches(inner.h))
    tf = box.text_frame
    tf.word_wrap = True
    hang = int(Pt(size) * 1.2)
    for i, b in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(size * 0.5)
        # 悬挂缩进：折行时第二行对齐到文字，而不是缩回圆点底下。python-pptx 没有
        # 现成的 API，直接写 a:pPr 的 marL / indent。圆点后面用制表符而不是空格：
        # 悬挂缩进下制表符正好跳到 marL，空格宽度随字体变，估不准就对不齐。
        ppr = p._p.get_or_add_pPr()
        ppr.set("marL", str(hang))
        ppr.set("indent", str(-hang))
        dot = p.add_run()
        dot.text = "•\t"
        _style_run(dot, size=size, color=BRAND_500, bold=True)
        r = p.add_run()
        r.text = b
        _style_run(r, size=size, color=TEXT_SECONDARY)


def _talk_picture(slide, photo: Path, rect: Rect, workdir: Path) -> bool:
    """在 rect 里等比放一张图，居中。任何异常都返回 False —— 配图挂了该页照出。"""
    from PIL import Image as _Image
    try:
        src = _embed_ready(photo, workdir)
        with _Image.open(src) as im:
            iw, ih = im.size
        scale = min(rect.w / iw, rect.h / ih)
        w, h = iw * scale, ih * scale
        pic = slide.shapes.add_picture(str(src), Inches(rect.x + (rect.w - w) / 2),
                                       Inches(rect.y), Inches(w), Inches(h))
        pic.line.color.rgb = RGBColor(*BORDER_SUBTLE)
        pic.line.width = Pt(0.75)
        return True
    except Exception as e:  # noqa: BLE001
        log.debug("讲稿配图插入失败：%s: %s", type(e).__name__, e)
        return False


def _talk_metrics(slide, rect: Rect, metrics: list[dict]) -> None:
    """规范里的指标卡：白底圆角卡片，大号数字用 brand-600，下面一行浅灰说明。"""
    n = max(1, len(metrics))
    gap = 0.3
    w = (rect.w - gap * (n - 1)) / n
    h = min(rect.h, 2.6)
    y = rect.y + (rect.h - h) / 2
    for i, m in enumerate(metrics):
        x = rect.x + i * (w + gap)
        _card(slide, Rect(x, y, w, h), radius=RADIUS_XL)
        value, label = str(m.get("value", "")), str(m.get("label", ""))
        _text(slide, Rect(x, y + 0.35, w, h * 0.5), value,
              size=fit_font_size(value, w, h * 0.5, max_pt=44), color=BRAND_600,
              bold=True, align=PP_ALIGN.CENTER)
        _text(slide, Rect(x, y + h * 0.58, w, h * 0.36), label,
              size=fit_font_size(label, w, h * 0.36, max_pt=16), color=TEXT_TERTIARY,
              align=PP_ALIGN.CENTER)


def _cell_borders(cell, bottom: RGB) -> None:
    """只留底边线（规范：td border-bottom 1px border-subtle），去掉竖线和上边线。

    默认表格样式会给每格画白色网格线；a:ln* 必须排在 tcPr 的最前面，顺序
    lnL → lnR → lnT → lnB，否则 PowerPoint 认为文件损坏。
    """
    tc_pr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        for old in tc_pr.findall(qn(tag)):
            tc_pr.remove(old)
    for i, tag in enumerate(("a:lnL", "a:lnR", "a:lnT", "a:lnB")):
        ln = etree.Element(qn(tag))
        if tag == "a:lnB":
            ln.set("w", "9525")
            fill = etree.SubElement(ln, qn("a:solidFill"))
            etree.SubElement(fill, qn("a:srgbClr")).set("val", "%02X%02X%02X" % bottom)
        else:
            ln.set("w", "0")
            etree.SubElement(ln, qn("a:noFill"))
        tc_pr.insert(i, ln)


TABLE_BODY_PT = 16
TABLE_HEAD_PT = 13
SHORT_COL_SHARE = 0.22        # 自然宽度不超过表宽这个比例的列，算「短列」


def _table_widths(grid: list[list[str]], ncols: int, total: float) -> list[float]:
    """列宽：短列（序号、日期、状态）先按自然宽度给足，不折行；剩下的宽度按内容
    长短分给长列。和浏览器 table-layout:auto 一个思路。

    只按内容长短比例分的话，「2026.07.31」这种列会被分到比它本身还窄的宽度，
    日期从中间断成两行。
    """
    natural = [max([_text_em(r[ci]) if ci < len(r) else 0.0 for r in grid] + [2.0])
               * TABLE_BODY_PT / 72 + 2 * CELL_PAD_X + 0.08
               for ci in range(ncols)]
    if sum(natural) <= total:
        return [total * n / sum(natural) for n in natural]
    short = [n <= total * SHORT_COL_SHARE for n in natural]
    fixed = sum(n for n, s in zip(natural, short) if s)
    rest = total - fixed
    longs = [n for n, s in zip(natural, short) if not s]
    if not longs or rest < total * 0.3:
        # 短列太多把长列挤没了：退回按比例分，每列保底最宽那列的四分之一。
        floor = max(natural) * 0.25
        weights = [max(floor, n) for n in natural]
        return [total * w / sum(weights) for w in weights]
    return [n if s else rest * n / sum(longs) for n, s in zip(natural, short)]


def _row_height(row: list[str], cell_w: list[float], pt: int) -> float:
    """一行在 pt 字号下要多高：最高那格的折行数 × 行距 + 上下内边距。"""
    lines = max([text_lines(t, cell_w[ci], pt) for ci, t in enumerate(row[:len(cell_w)])] or [1])
    return lines * pt * LINE_SPACING / 72 + 2 * CELL_PAD_Y + 0.04


def _talk_table(slide, rect: Rect, columns: list[str], rows: list[list[str]]) -> None:
    """原生表格 —— 要能在 PowerPoint 里直接改。

    样式按规范：外层白色圆角卡片（原生表格本身做不了圆角），表头浅灰底 + 等宽字 +
    三级灰字，正文白底 + 二级灰字，行与行之间一条淡分隔线，没有竖线。
    """
    ncols = max([len(columns)] + [len(r) for r in rows]) if (columns or rows) else 1
    nrows = len(rows) + (1 if columns else 0)
    pad = 0.12
    inner = Rect(rect.x + pad, rect.y + pad, rect.w - 2 * pad, rect.h - 2 * pad)
    grid = ([columns] if columns else []) + rows
    widths = _table_widths(grid, ncols, inner.w)
    cell_w = [w - 2 * CELL_PAD_X + TEXT_MARGIN_X for w in widths]
    # 行高按内容、全表一个字号（表头另一个）：从大往小试，取整张表装得下的最大字号。
    # 以前每行等高、封顶 0.75in，一格长文字就把全表压到很小，而页面下半截空着。
    head_pt = min([fit_font_size(c, cell_w[ci], 0.75, max_pt=TABLE_HEAD_PT)
                   for ci, c in enumerate(columns[:ncols])] or [TABLE_HEAD_PT])
    heights: list[float] = []
    for body_pt in range(TABLE_BODY_PT, MIN_FONT_PT - 1, -1):
        heights = [_row_height(r, cell_w, head_pt if (columns and ri == 0) else body_pt)
                   for ri, r in enumerate(grid)]
        if sum(heights) <= inner.h:
            break
    table_h = sum(heights)
    _card(slide, Rect(rect.x, rect.y, rect.w, min(rect.h, table_h + 2 * pad)), radius=RADIUS_XL)
    shape = slide.shapes.add_table(nrows, ncols, Inches(inner.x), Inches(inner.y),
                                   Inches(inner.w), Inches(table_h))
    table = shape.table
    table.horz_banding = False
    table.first_row = bool(columns)
    for ci, w in enumerate(widths):
        table.columns[ci].width = Emu(int(Inches(w)))
    for ri, h in enumerate(heights):
        table.rows[ri].height = Emu(int(Inches(h)))
    for ri, row in enumerate(grid):
        header = bool(columns) and ri == 0
        last = ri == len(grid) - 1
        for ci in range(ncols):
            text = row[ci] if ci < len(row) else ""
            cell = table.cell(ri, ci)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(*(SURFACE_SUBTLE if header else SURFACE_ELEVATED))
            _cell_borders(cell, BORDER_DEFAULT if header else (SURFACE_ELEVATED if last else BORDER_SUBTLE))
            cell.margin_left = cell.margin_right = Emu(int(Inches(CELL_PAD_X)))
            cell.margin_top = cell.margin_bottom = Emu(int(Inches(CELL_PAD_Y)))
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE     # 行高给足之后，贴顶的字会显得不平衡
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = ""
            r = p.add_run()
            r.text = text
            _style_run(r, size=head_pt if header else body_pt,
                       color=TEXT_TERTIARY if header else TEXT_SECONDARY,
                       bold=header, mono=header)


def render_talk(talk: Talk, images: dict[str, Path], appendix: list[str], out_path: Path,
                appendix_notes: str = "") -> Path:
    """讲稿 → .pptx：封面 + 各页 + （有未收录章节时）末尾一页「未收录章节」。

    配色按 Lumen-light：封面是品牌绿渐变 Banner，正文页浅底 + 白卡片，和 HTML 产物一套。
    """
    import tempfile

    prs = Presentation()
    prs.slide_width = Emu(int(Inches(SLIDE_W)))
    prs.slide_height = Emu(int(Inches(SLIDE_H)))
    workdir = Path(tempfile.mkdtemp(prefix="lah-talk-"))
    body_y = 0.3 + TALK_TITLE_H + 0.25
    body_h = SLIDE_H - body_y - MARGIN
    body_w = SLIDE_W - 2 * MARGIN

    _talk_cover(prs, talk.title, talk.subtitle)

    for s in talk.slides:
        slide = _new_talk_slide(prs)
        _talk_title(slide, s.title)
        body = Rect(MARGIN, body_y, body_w, body_h)
        if s.kind == "metrics":
            _talk_metrics(slide, body, s.metrics)
        elif s.kind == "table":
            _talk_table(slide, body, s.columns, s.rows)
        else:
            photo = images.get(s.image) if s.image else None
            if photo is not None:
                text_w = body_w * (1 - IMAGE_SHARE) - 0.3
                _bullet_box(slide, Rect(MARGIN, body_y, text_w, body_h), s.bullets)
                _talk_picture(slide, photo,
                              Rect(MARGIN + text_w + 0.3, body_y, body_w * IMAGE_SHARE, body_h), workdir)
            else:
                _bullet_box(slide, body, s.bullets)
        if s.notes:
            slide.notes_slide.notes_text_frame.text = s.notes

    if appendix:
        slide = _new_talk_slide(prs)
        _talk_title(slide, "未收录章节")
        _text(slide, Rect(MARGIN, body_y - 0.1, body_w, 0.45),
              "以下章节在讲稿里没有用上，需要的话手动补回。", size=14, color=TEXT_TERTIARY)
        _bullet_box(slide, Rect(MARGIN, body_y + 0.45, body_w, body_h - 0.45), appendix)
        if appendix_notes:
            slide.notes_slide.notes_text_frame.text = appendix_notes

    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    shutil.rmtree(workdir, ignore_errors=True)
    return out_path
