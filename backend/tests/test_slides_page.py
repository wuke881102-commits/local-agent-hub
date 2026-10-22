import pytest

from app.agents.base import AgentContext
from app.agents.slides_page import SlidesPageAgent, pick_mode
from app.routes.tasks import _download_meta


# ── 模式判定（阈值取自实测的四篇文档）──────────────────────────────

def test_pick_mode_image_heavy():
    """实测：图片型 58 / 59 字每张。"""
    assert pick_mode(chars=582, images=10) == "restore"
    assert pick_mode(chars=297, images=5) == "restore"


def test_pick_mode_text_heavy():
    """实测：文字型 509 / 4741 字每张。"""
    assert pick_mode(chars=28534, images=56) == "outline"
    assert pick_mode(chars=9483, images=2) == "outline"


def test_pick_mode_too_few_images():
    """图少于 3 张，再稀疏也不走还原 —— 还原一两张图不值当一份 deck。"""
    assert pick_mode(chars=10, images=2) == "outline"


def test_pick_mode_no_images():
    assert pick_mode(chars=1000, images=0) == "outline"


async def test_agent_fails_without_token():
    r = await SlidesPageAgent().run(
        AgentContext(task_id="t", agent_id="slides-page", inputs={}, lark=None, llm=None))
    assert r.status == "failed"
    assert "doc_token" in (r.error or "")


async def test_agent_rejects_doc_without_images(monkeypatch):
    from app.services import doc_source

    async def fake_load(ctx, token, inputs):
        return {"markdown": "只有文字没有图", "title": "无图"}

    monkeypatch.setattr(doc_source, "load_source", fake_load)
    r = await SlidesPageAgent().run(
        AgentContext(task_id="t", agent_id="slides-page",
                     inputs={"doc_token": "T", "ppt_mode": "restore"},
                     lark=None, llm=None))
    assert r.status == "failed"
    assert "没有可还原的图片" in (r.error or "")


# ── 下载路由 ──────────────────────────────────────────────────────

def test_download_meta_html():
    assert _download_meta("t-1", "/x/t-1.html") == ("text/html", "t-1.html")


def test_download_meta_pptx():
    media, name = _download_meta("t-1", "/x/t-1.pptx")
    assert media == ("application/vnd.openxmlformats-officedocument"
                     ".presentationml.presentation")
    assert name == "t-1.pptx"


def test_download_meta_unknown_extension_falls_back_to_octet_stream():
    """认不出来就给 octet-stream —— 让浏览器下载，而不是猜错类型当网页渲染。"""
    assert _download_meta("t-1", "/x/t-1.bin") == ("application/octet-stream", "t-1.bin")


async def test_preview_refuses_non_html_results(tmp_path, monkeypatch):
    """/preview 只认 HTML。原先对任何产物都 read_text(utf-8)，.pptx 是二进制，
    走到这里就是 UnicodeDecodeError → 500，前端只看到一块白。"""
    from fastapi import HTTPException
    from app.routes import tasks as tasks_route

    f = tmp_path / "t-1.pptx"
    f.write_bytes(b"PK\x03\x04\xff\xfe binary")

    async def fake_get(task_id):
        return {"id": task_id, "result_path": str(f)}

    monkeypatch.setattr(tasks_route.task_runner, "get_task", fake_get)
    try:
        await tasks_route.preview_task("t-1")
        assert False, "应当拒绝"
    except HTTPException as e:
        assert e.status_code == 415
        assert "下载" in e.detail


async def test_preview_still_serves_html(tmp_path, monkeypatch):
    from app.routes import tasks as tasks_route

    f = tmp_path / "t-2.html"
    f.write_text("<h1>页面</h1>", encoding="utf-8")

    async def fake_get(task_id):
        return {"id": task_id, "result_path": str(f)}

    monkeypatch.setattr(tasks_route.task_runner, "get_task", fake_get)
    r = await tasks_route.preview_task("t-2")
    assert "页面" in r.body.decode("utf-8")


async def test_download_failure_reports_the_real_reason(monkeypatch):
    """图片下载失败时，日志和任务错误里要有**真实原因**，不能只有异常类名。

    实测：开发后端用的全局 CLI 是 1.0.89，media-download 被一个 scope 预检拒掉，
    而页面上只看得到十行「下载失败，已跳过：LarkCLIError」和一句
    「所有图片都没能处理成功」—— 8.2 在 doc_images 里修过的同一个毛病，
    在 slides_page 里又写了一遍。
    """
    from app.feishu.cli import LarkCLIError
    from app.services import doc_source

    async def fake_load(ctx, token, inputs):
        imgs = "".join(f"![](https://feishu.cn/file/TOKEN{i:022d})" for i in range(4))
        return {"markdown": "标题" + imgs, "title": "图片型"}

    class Lark:
        async def docs_download_media(self, tok, d, name):
            raise LarkCLIError("lark-cli exited 1", exit_code=1, stderr=(
                '{"ok":false,"error":{"type":"authorization","subtype":"missing_scope",'
                '"message":"missing required scope(s): docs:permission.member:auth"}}'))

    monkeypatch.setattr(doc_source, "load_source", fake_load)
    lines = []

    async def emit(level, msg):
        lines.append(msg)

    r = await SlidesPageAgent().run(AgentContext(
        task_id="t", agent_id="slides-page", inputs={"doc_token": "T"},
        lark=Lark(), llm=None, emit=emit))
    assert r.status == "failed"
    assert "LarkCLIError" not in (r.error or ""), r.error
    assert "缺少飞书权限" in (r.error or ""), f"任务错误里没有真实原因：{r.error}"
    assert "4/4" in (r.error or "")
    assert not any(x.endswith("LarkCLIError") for x in lines), "日志里仍然只有异常类名"
    # 同一个原因只完整说一次，后面几张不刷屏
    assert sum("缺少飞书权限" in x for x in lines) == 1


# ── 讲稿模式（文字型文档）────────────────────────────────────────

import json as _json

TEXT_DOC = (
    "# 产品使用培训\n\n"
    + "".join(f"## 第{i}节\n" + f"第{i}节的正文，讲一件具体的事情。" * 12 + "\n\n" for i in range(1, 5))
    + "![界面截图，左侧是消息列表](https://feishu.cn/file/AAAAAAAAAAAAAAAAAAAAAAAA)\n"
)


def _talk_json(slides):
    return _json.dumps({"title": "讲稿", "subtitle": "副标题", "slides": slides}, ensure_ascii=False)


def _bs(title, sources, **kw):
    return {"kind": "bullets", "title": title, "bullets": ["要点一", "要点二"], "sources": sources, **kw}


class TalkLLM:
    vision_model = "fake"

    def __init__(self, replies):
        self.replies, self.calls = list(replies), 0

    async def text_complete(self, prompt, system=None, **kw):
        self.calls += 1
        return self.replies.pop(0) if self.replies else ""


class NoLark:
    async def docs_download_media(self, tok, d, name):
        raise AssertionError("不该下载")


@pytest.fixture
def text_doc(monkeypatch):
    from app.services import doc_source

    async def fake_load(ctx, token, inputs):
        return {"markdown": TEXT_DOC, "title": "产品使用培训"}

    monkeypatch.setattr(doc_source, "load_source", fake_load)


async def _run(llm, lark=None):
    lines = []

    async def emit(level, msg):
        lines.append(msg)

    r = await SlidesPageAgent().run(AgentContext(
        task_id="t-talk", agent_id="slides-page", inputs={"doc_token": "T"},
        lark=lark or NoLark(), llm=llm, emit=emit))
    return r, lines


async def test_text_doc_becomes_a_talk_deck(text_doc, tmp_path, monkeypatch):
    from app.config import settings
    monkeypatch.setattr(settings, "draft_dir", str(tmp_path))
    llm = TalkLLM([_talk_json([_bs("一", ["S1"]), _bs("二", ["S2"]), _bs("三", ["S3", "S4"])])])
    r, lines = await _run(llm)
    assert r.status == "preview", r.error
    assert r.result_path.endswith(".pptx")
    assert r.payload["mode"] == "outline"
    assert r.payload["slides"] == 1 + 3                 # 封面 + 3 页，全覆盖所以没有附录
    assert r.payload["uncovered"] == 0
    assert any("讲稿" in x for x in lines)


async def test_uncovered_sections_are_logged_and_listed(text_doc, tmp_path, monkeypatch):
    from app.config import settings
    from pptx import Presentation
    monkeypatch.setattr(settings, "draft_dir", str(tmp_path))
    llm = TalkLLM([_talk_json([_bs("一", ["S1"]), _bs("二", ["S1"]), _bs("三", ["S2"])])])
    r, lines = await _run(llm)
    assert r.payload["uncovered"] == 2
    assert any("第3节" in x and "第4节" in x for x in lines), "日志里要点名没用上的章节"
    last = Presentation(r.result_path).slides[-1]
    text = "\n".join(sh.text_frame.text for sh in last.shapes if sh.has_text_frame)
    assert "未收录章节" in text and "第3节" in text and "第4节" in text


async def test_garbage_twice_fails_with_the_reason(text_doc):
    """模型两次都给不出能解析的讲稿，就明确失败 —— 不交一份空 deck。"""
    llm = TalkLLM(["抱歉我做不到", "还是做不到"])
    r, _ = await _run(llm)
    assert r.status == "failed"
    assert llm.calls == 2
    assert "讲稿生成失败" in (r.error or "") and "不是 JSON" in (r.error or "")


async def test_retry_succeeds_on_second_attempt(text_doc, tmp_path, monkeypatch):
    from app.config import settings
    monkeypatch.setattr(settings, "draft_dir", str(tmp_path))
    ok = _talk_json([_bs("一", ["S1", "S2"]), _bs("二", ["S3"]), _bs("三", ["S4"])])
    llm = TalkLLM(["不是 json", ok])
    r, _ = await _run(llm)
    assert r.status == "preview" and llm.calls == 2


async def test_image_download_failure_keeps_the_deck(text_doc, tmp_path, monkeypatch):
    """配图挂了该页照出不带图，原因写进日志，任务仍然成功。"""
    from app.config import settings
    from app.feishu.cli import LarkCLIError
    monkeypatch.setattr(settings, "draft_dir", str(tmp_path))

    class BadLark:
        async def docs_download_media(self, tok, d, name):
            raise LarkCLIError("lark-cli exited 1", exit_code=1, stderr=(
                '{"ok":false,"error":{"type":"authorization","subtype":"missing_scope",'
                '"message":"missing required scope(s): docs:permission.member:auth"}}'))

    llm = TalkLLM([_talk_json([_bs("一", ["S1"], image="I1"), _bs("二", ["S2"]),
                               _bs("三", ["S3", "S4"])])])
    r, lines = await _run(llm, BadLark())
    assert r.status == "preview"
    assert r.payload["images"] == 0
    assert any("缺少飞书权限" in x for x in lines)
    assert not any(x.endswith("LarkCLIError") for x in lines)


async def test_dropped_items_are_logged(text_doc, tmp_path, monkeypatch):
    """校验时丢掉的东西（编造的数字、引用不存在的图）要进日志，不能静默。"""
    from app.config import settings
    monkeypatch.setattr(settings, "draft_dir", str(tmp_path))
    llm = TalkLLM([_talk_json([_bs("一", ["S1"], image="I9"), _bs("二", ["S2"]),
                               _bs("三", ["S3", "S4"])])])
    r, lines = await _run(llm)
    assert r.status == "preview"
    assert any("I9" in x for x in lines)
