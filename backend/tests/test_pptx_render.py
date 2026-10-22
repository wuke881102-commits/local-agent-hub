from pathlib import Path

from PIL import Image
from pptx.oxml.ns import qn

from app.services.pptx_render import Block, Column, Band, SlidePlan


def test_slideplan_defaults():
    p = SlidePlan(title="标题")
    assert p.title == "标题"
    assert p.theme == "dark"
    assert p.bands == []
    assert p.is_slide is True


def test_block_group_items_are_independent():
    """dataclass 的可变默认值不能共享 —— 共享会让两个 block 的 items 互相污染。"""
    a, b = Block(kind="group"), Block(kind="group")
    a.items.append("x")
    assert b.items == []


from app.services.pptx_render import Rect, band_heights, column_rects


def test_band_heights_weighted_by_content():
    """两条带，一条 3 块一条 1 块 —— 高度按内容量加权，不是均分。"""
    thick = Band(columns=[Column(blocks=[Block(), Block(), Block()])])
    thin = Band(columns=[Column(blocks=[Block()])])
    hs = band_heights([thick, thin], total=8.0, gap=0.2)
    assert len(hs) == 2
    assert hs[0] > hs[1]
    assert abs(sum(hs) + 0.2 - 8.0) < 1e-6      # 高度 + 带间距 = 总高


def test_band_heights_single_band_takes_all():
    hs = band_heights([Band(columns=[Column(blocks=[Block()])])], total=8.0, gap=0.2)
    assert abs(hs[0] - 8.0) < 1e-6


def test_band_heights_empty_bands_never_divides_by_zero():
    """没有块的带也要有高度，否则会出现 0 高的形状（python-pptx 会报错）。"""
    hs = band_heights([Band(), Band()], total=8.0, gap=0.2)
    assert all(h > 0 for h in hs)


def test_band_heights_no_bands():
    assert band_heights([], total=8.0, gap=0.2) == []


def test_column_rects_split_width_evenly():
    band = Band(columns=[Column(), Column(), Column(), Column()])
    rects = column_rects(band, Rect(0.5, 1.0, 12.0, 5.0), gap=0.2)
    assert len(rects) == 4
    assert abs(rects[0].x - 0.5) < 1e-6
    assert all(abs(r.w - rects[0].w) < 1e-6 for r in rects)
    assert rects[-1].x + rects[-1].w <= 0.5 + 12.0 + 1e-6


def test_column_rects_no_columns_returns_empty():
    assert column_rects(Band(), Rect(0, 0, 10, 5), gap=0.2) == []


from app.services.pptx_render import MIN_FONT_PT, fit_font_size


def test_fit_font_size_shrinks_for_long_text():
    short = fit_font_size("短", width=2.0, height=0.6)
    long = fit_font_size("这是一段相当长的说明文字需要塞进同样大小的框里", width=2.0, height=0.6)
    assert short > long


def test_fit_font_size_never_below_floor():
    """再长也不低于下限 —— 低于 9pt 投影上没人看得见，不如截断。"""
    assert fit_font_size("字" * 500, width=1.0, height=0.3) == MIN_FONT_PT


def test_fit_font_size_handles_empty_text():
    assert fit_font_size("", width=2.0, height=0.6) > 0


from pptx import Presentation

from app.services.pptx_render import DeckItem, render_deck


def _plan() -> SlidePlan:
    return SlidePlan(
        title="双通道协同",
        subtitle="统一获取多类代码源",
        bands=[Band(role="columns", columns=[
            Column(index="00", heading="数据源接入", blocks=[
                Block(kind="box", label="统一数据快照", sub="仓库 / 分支 / Commit"),
                Block(kind="group", label="多入口", items=["GitLab", "Jenkins"]),
            ]),
            Column(index="01", heading="上下文资产化", blocks=[
                Block(kind="note", label="一次解析，多引擎复用"),
            ]),
        ])],
        notes="全文转写内容",
    )


def _photo(tmp_path: Path) -> Path:
    f = tmp_path / "p.png"
    Image.new("RGB", (400, 300), (12, 20, 40)).save(f)
    return f


def test_render_deck_makes_two_slides_per_item(tmp_path: Path):
    """一张图 → 一页还原 + 一页原图。"""
    out = render_deck([DeckItem(plan=_plan(), photo=_photo(tmp_path))], tmp_path / "d.pptx")
    assert len(Presentation(str(out)).slides) == 2


def test_render_deck_writes_all_labels(tmp_path: Path):
    out = render_deck([DeckItem(plan=_plan(), photo=_photo(tmp_path))], tmp_path / "d.pptx")
    prs = Presentation(str(out))
    text = "\n".join(sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame)
    for expected in ("双通道协同", "00", "数据源接入", "统一数据快照",
                     "仓库 / 分支 / Commit", "GitLab", "Jenkins", "一次解析，多引擎复用"):
        assert expected in text, f"缺少 {expected!r}"


def test_render_deck_puts_transcript_in_notes(tmp_path: Path):
    out = render_deck([DeckItem(plan=_plan(), photo=_photo(tmp_path))], tmp_path / "d.pptx")
    prs = Presentation(str(out))
    assert "全文转写内容" in prs.slides[0].notes_slide.notes_text_frame.text


def test_render_deck_non_slide_photo_only(tmp_path: Path):
    """is_slide=False（合影之类）只出一页原图，不硬还原成一堆框。"""
    plan = SlidePlan(is_slide=False, title="现场合影", notes="一张合影")
    out = render_deck([DeckItem(plan=plan, photo=_photo(tmp_path))], tmp_path / "d.pptx")
    assert len(Presentation(str(out)).slides) == 1


def test_render_deck_without_photo_still_renders(tmp_path: Path):
    """图下载失败时仍要出还原页，不能整份 deck 报废。"""
    out = render_deck([DeckItem(plan=_plan(), photo=None)], tmp_path / "d.pptx")
    assert len(Presentation(str(out)).slides) == 1


def test_render_deck_empty_items_makes_valid_file(tmp_path: Path):
    out = render_deck([], tmp_path / "d.pptx")
    assert out.exists()
    assert len(Presentation(str(out)).slides) == 0


def test_render_deck_is_widescreen(tmp_path: Path):
    from pptx.util import Inches
    out = render_deck([DeckItem(plan=_plan(), photo=None)], tmp_path / "d.pptx")
    prs = Presentation(str(out))
    assert abs(prs.slide_width / Inches(1) - 13.333) < 0.01


def test_render_deck_callout_keeps_its_label(tmp_path: Path):
    """callout 的 value 是大字、label 是它的说明，两个都要画出来。

    这是肉眼看渲染结果才发现的：value 占了主文字位，label 没有任何去处，
    「↑80%」画出来了而「客服响应效率」凭空消失 —— 20 条断言全绿也没拦住。
    """
    plan = SlidePlan(title="t", bands=[Band(columns=[Column(blocks=[
        Block(kind="callout", value="↑80%", label="客服响应效率"),
    ])])])
    out = render_deck([DeckItem(plan=plan, photo=None)], tmp_path / "d.pptx")
    text = "\n".join(sh.text_frame.text for sh in Presentation(str(out)).slides[0].shapes
                     if sh.has_text_frame)
    assert "↑80%" in text
    assert "客服响应效率" in text


# ── 真文档跑出来才暴露的三个问题 ──────────────────────────────────

from app.services.pptx_render import MIN_BLOCK_H, stack_rects


def test_stack_rects_stay_inside_area():
    """块再多也不许溢出给定区域 —— 溢出就会压到下一条带上。"""
    area = Rect(0.5, 1.0, 3.0, 2.0)
    for r in stack_rects(area, n=12, gap=0.16, min_h=MIN_BLOCK_H):
        assert r.y >= area.y - 1e-6
        assert r.y + r.h <= area.y + area.h + 1e-6


def test_stack_rects_drops_blocks_that_cannot_fit():
    """放不下就少画几个（其余进备注），而不是硬塞到越界。"""
    rects = stack_rects(Rect(0, 0, 3.0, 0.9), n=12, gap=0.16, min_h=MIN_BLOCK_H)
    assert 0 < len(rects) < 12


def test_stack_rects_never_returns_zero_height():
    for r in stack_rects(Rect(0, 0, 3.0, 0.2), n=5, gap=0.16, min_h=MIN_BLOCK_H):
        assert r.h > 0


def test_stack_rects_all_fit_when_there_is_room():
    assert len(stack_rects(Rect(0, 0, 3.0, 6.0), n=4, gap=0.16, min_h=MIN_BLOCK_H)) == 4


def test_render_deck_bands_do_not_overlap(tmp_path: Path):
    """两条带，第一条塞 10 块 —— 第一条的形状不许伸进第二条的地盘。

    这是真文档跑出来才看见的：块高有个 0.3 英寸下限，块一多堆叠总高就超过
    本带高度，直接压到下一带上，两带文字叠在一起。
    """
    from pptx.util import Inches
    plan = SlidePlan(title="t", bands=[
        Band(columns=[Column(blocks=[Block(label=f"块{i}") for i in range(10)])]),
        Band(columns=[Column(blocks=[Block(label="第二带")])]),
    ])
    out = render_deck([DeckItem(plan=plan, photo=None)], tmp_path / "d.pptx")
    prs = Presentation(str(out))
    slide = prs.slides[0]
    bottoms = [sh.top + sh.height for sh in slide.shapes]
    assert max(bottoms) <= prs.slide_height, "有形状画到了页面之外"


def test_render_deck_downscales_embedded_photos(tmp_path: Path):
    """原图页是拿来核对的，不需要 4032px 的手机原图。

    实测 10 张图的 deck 有 37 MB，22 张就是 80 MB —— 发不出去。
    缩到长边 PHOTO_MAX_PX 之后屏幕上照样看得清每个字。
    """
    from app.services.pptx_render import PHOTO_MAX_PX

    big = tmp_path / "big.jpg"
    Image.new("RGB", (4032, 3024), (30, 40, 60)).save(big, quality=95)
    out = render_deck([DeckItem(plan=SlidePlan(is_slide=False, title="t"), photo=big)],
                      tmp_path / "d.pptx")

    import zipfile
    with zipfile.ZipFile(out) as z:
        media = [n for n in z.namelist() if n.startswith("ppt/media/")]
        assert media, "没有嵌入图片"
        with z.open(media[0]) as f:
            w, h = Image.open(f).size
    assert max(w, h) <= PHOTO_MAX_PX, f"嵌入的还是 {w}x{h}"


def test_render_deck_keeps_small_photos_untouched(tmp_path: Path):
    """本来就不大的图不要再压一道，白白掉画质。"""
    from app.services.pptx_render import PHOTO_MAX_PX

    small = tmp_path / "small.png"
    Image.new("RGB", (800, 600), (30, 40, 60)).save(small)
    out = render_deck([DeckItem(plan=SlidePlan(is_slide=False), photo=small)],
                      tmp_path / "d.pptx")
    import zipfile
    with zipfile.ZipFile(out) as z:
        media = [n for n in z.namelist() if n.startswith("ppt/media/")]
        with z.open(media[0]) as f:
            w, h = Image.open(f).size
    assert (w, h) == (800, 600)
    assert max(w, h) <= PHOTO_MAX_PX


# ── 8.3 装机版上肉眼看到的：长文字溢出压到栏标题 ────────────────────

from app.services.pptx_render import text_lines


def test_fit_font_size_actually_fits_the_box():
    """选出来的字号，按换行后的行数算，必须装得进框。

    8.3 的算法是「面积 ÷ 字数」，没算形状内边距、行距、换行，于是估大了；
    PowerPoint 对装不下的文字上下对称溢出，往上就压到了栏标题。
    """
    from app.services.pptx_render import LINE_SPACING, TEXT_MARGIN_Y, fit_font_size
    text = "审核Agent重新读取代码与证据，不直接采用引擎结论"
    w, h = 3.0, 0.75
    pt = fit_font_size(text, w, h)
    need = text_lines(text, w, pt) * pt * LINE_SPACING / 72 + TEXT_MARGIN_Y
    assert need <= h + 1e-6 or pt == MIN_FONT_PT, f"{pt}pt 需要 {need:.2f}\" 但框只有 {h}\""


def test_text_lines_counts_wrapping_and_paragraphs():
    assert text_lines("短", 3.0, 12) == 1
    assert text_lines("第一段\n第二段", 3.0, 12) == 2
    long = "字" * 60
    assert text_lines(long, 3.0, 12) > text_lines(long, 3.0, 9) >= 1


def test_text_lines_latin_is_narrower_than_cjk():
    """英文字母约半个中文字宽，否则英文多的框会被估得太挤、字号给得太小。"""
    assert text_lines("a" * 40, 2.0, 12) < text_lines("字" * 40, 2.0, 12)


def test_stack_rects_gives_long_blocks_more_height():
    """一栏里一个长块加几个短块：长块应该更高，而不是大家一样高。"""
    rects = stack_rects(Rect(0, 0, 3.0, 6.0), n=4, gap=0.16, min_h=MIN_BLOCK_H,
                        weights=[5.0, 1.0, 1.0, 1.0])
    assert len(rects) == 4
    assert rects[0].h > rects[1].h * 2
    assert abs(rects[1].h - rects[2].h) < 1e-6
    assert rects[-1].y + rects[-1].h <= 6.0 + 1e-6


def test_stack_rects_weights_never_starve_a_block_below_min():
    rects = stack_rects(Rect(0, 0, 3.0, 3.0), n=4, gap=0.1, min_h=MIN_BLOCK_H,
                        weights=[100.0, 1.0, 1.0, 1.0])
    assert all(r.h >= MIN_BLOCK_H - 1e-6 for r in rects)


def test_render_deck_no_block_text_overflows(tmp_path: Path):
    """复现 8.3 那一栏：一个长块 + 若干短块。每个框按选定字号都要装得下。"""
    from pptx.util import Emu
    from app.services.pptx_render import LINE_SPACING, TEXT_MARGIN_Y
    long = "审核Agent重新读取代码与证据，不直接采用引擎结论"
    col = Column(index="03", heading="审核与可复核输出", blocks=(
        [Block(label=long)] + [Block(label=t) for t in
         ("候选风险集合", "证据一致性：代码、变量、文件位置相互印证",
          "路径可达：Source到Sink真实连通", "可利用性：净化有效性与PoC",
          "高置信风险", "类型与优先级")]))
    plan = SlidePlan(title="t", bands=[Band(columns=[Column(), Column(), Column(), col])])
    out = render_deck([DeckItem(plan=plan, photo=None)], tmp_path / "d.pptx")
    for sh in Presentation(str(out)).slides[0].shapes:
        if not sh.has_text_frame or sh.shape_type != 1:     # 只查自选图形（框）
            continue
        runs = [r for p in sh.text_frame.paragraphs for r in p.runs]
        if not runs:
            continue
        pt = runs[0].font.size.pt
        w, h = Emu(sh.width).inches, Emu(sh.height).inches
        need = text_lines(sh.text_frame.text, w, pt) * pt * LINE_SPACING / 72 + TEXT_MARGIN_Y
        assert need <= h + 0.02 or pt <= MIN_FONT_PT, \
            f"「{sh.text_frame.text[:14]}」{pt}pt 需要 {need:.2f}\"，框只有 {h:.2f}\""


# ── 讲稿模式渲染 ──────────────────────────────────────────────────

from app.services.pptx_render import Talk, TalkSlide, render_talk


def _deck(tmp_path, **kw) -> Talk:
    return Talk(title="产品能力培训", subtitle="从开通到日常使用", slides=[
        TalkSlide(kind="bullets", title="能做什么", bullets=["写文档", "做表格", "查资料"],
                  notes="原文细节一", sources=["S1"], image=kw.get("image")),
        TalkSlide(kind="metrics", title="关键数据",
                  metrics=[{"value": "80%", "label": "效率提升"}, {"value": "1200 万", "label": "拦截次数"}],
                  sources=["S2"]),
        TalkSlide(kind="table", title="分工", columns=["阶段", "负责人"],
                  rows=[["开发", "张三"], ["测试", "李四"]], sources=["S3"]),
        TalkSlide(kind="closing", title="下一步", bullets=["开通账号", "参加第二期培训"], sources=["S4"]),
    ])


def _all_text(slide) -> str:
    out = []
    for sh in slide.shapes:
        if sh.has_text_frame:
            out.append(sh.text_frame.text)
        if sh.has_table:
            out.extend(c.text for r in sh.table.rows for c in r.cells)
    return "\n".join(out)


def test_render_talk_page_count(tmp_path: Path):
    out = render_talk(_deck(tmp_path), {}, [], tmp_path / "t.pptx")
    assert len(Presentation(str(out)).slides) == 1 + 4          # 封面 + 4 页


def test_render_talk_appendix_page_when_sections_uncovered(tmp_path: Path):
    out = render_talk(_deck(tmp_path), {}, ["二、怎么开通", "五、附录"], tmp_path / "t.pptx")
    prs = Presentation(str(out))
    assert len(prs.slides) == 1 + 4 + 1
    last = _all_text(prs.slides[-1])
    assert "未收录章节" in last and "二、怎么开通" in last and "五、附录" in last


def test_render_talk_cover_has_title_and_subtitle(tmp_path: Path):
    out = render_talk(_deck(tmp_path), {}, [], tmp_path / "t.pptx")
    t = _all_text(Presentation(str(out)).slides[0])
    assert "产品能力培训" in t and "从开通到日常使用" in t


def test_render_talk_each_kind_shows_its_content(tmp_path: Path):
    out = render_talk(_deck(tmp_path), {}, [], tmp_path / "t.pptx")
    s = Presentation(str(out)).slides
    assert all(x in _all_text(s[1]) for x in ("能做什么", "写文档", "做表格", "查资料"))
    assert all(x in _all_text(s[2]) for x in ("80%", "效率提升", "1200 万", "拦截次数"))
    assert all(x in _all_text(s[3]) for x in ("阶段", "负责人", "开发", "张三", "李四"))
    assert all(x in _all_text(s[4]) for x in ("下一步", "开通账号", "参加第二期培训"))


def test_render_talk_table_is_a_native_table(tmp_path: Path):
    """表格要能在 PowerPoint 里直接改，不能是一张图或一堆文本框。"""
    out = render_talk(_deck(tmp_path), {}, [], tmp_path / "t.pptx")
    slide = Presentation(str(out)).slides[3]
    tables = [sh for sh in slide.shapes if sh.has_table]
    assert len(tables) == 1
    assert len(tables[0].table.rows) == 3 and len(tables[0].table.columns) == 2


def test_render_talk_writes_notes(tmp_path: Path):
    out = render_talk(_deck(tmp_path), {}, [], tmp_path / "t.pptx")
    assert "原文细节一" in Presentation(str(out)).slides[1].notes_slide.notes_text_frame.text


def test_render_talk_embeds_the_chosen_image(tmp_path: Path):
    img = tmp_path / "i1.png"
    Image.new("RGB", (800, 600), (200, 100, 50)).save(img)
    out = render_talk(_deck(tmp_path, image="I1"), {"I1": img}, [], tmp_path / "t.pptx")
    pics = [sh for sh in Presentation(str(out)).slides[1].shapes if sh.shape_type == 13]
    assert len(pics) == 1


def test_render_talk_survives_a_missing_image(tmp_path: Path):
    """配图下载失败时该页照出、不带图，整份 deck 不能因此报废。"""
    out = render_talk(_deck(tmp_path, image="I1"), {}, [], tmp_path / "t.pptx")
    pics = [sh for sh in Presentation(str(out)).slides[1].shapes if sh.shape_type == 13]
    assert pics == []


def test_render_talk_bullet_text_fits(tmp_path: Path):
    """要点区按选定字号要装得下 —— 按图还原那边溢出压到标题的教训。"""
    from pptx.util import Emu
    from app.services.pptx_render import LINE_SPACING, TEXT_MARGIN_Y
    long = ["这是一条相当长的要点，用来测试要点区的字号是否会自动缩小到装得下为止" * 2] * 6
    talk = Talk(title="t", slides=[TalkSlide(kind="bullets", title="长要点", bullets=long)] * 3)
    out = render_talk(talk, {}, [], tmp_path / "t.pptx")
    for sh in Presentation(str(out)).slides[1].shapes:
        if not sh.has_text_frame or "相当长" not in sh.text_frame.text:
            continue
        pt = sh.text_frame.paragraphs[0].runs[0].font.size.pt
        need = text_lines(sh.text_frame.text, Emu(sh.width).inches, pt) * pt * LINE_SPACING / 72 + TEXT_MARGIN_Y
        assert need <= Emu(sh.height).inches + 0.02 or pt <= MIN_FONT_PT


def test_render_talk_is_widescreen(tmp_path: Path):
    from pptx.util import Inches
    out = render_talk(_deck(tmp_path), {}, [], tmp_path / "t.pptx")
    assert abs(Presentation(str(out)).slide_width / Inches(1) - 13.333) < 0.01


def test_render_talk_table_gives_wide_content_more_width(tmp_path: Path):
    """列宽按内容分：「序号」这种一两个字的列不该和「状态更新」一样宽。"""
    talk = Talk(title="t", slides=[TalkSlide(kind="table", title="进展", columns=["序号", "状态更新"],
                rows=[["1", "调研准备阶段完成，项目实施阶段进行中，计划十月输出报告"]])] * 3)
    out = render_talk(talk, {}, [], tmp_path / "t.pptx")
    tb = next(sh for sh in Presentation(str(out)).slides[1].shapes if sh.has_table).table
    assert tb.columns[1].width > tb.columns[0].width * 3


# ── Lumen-light 样式 ────────────────────────────────────────────────

def _xml(prs) -> str:
    return "\n".join(sl._element.xml for sl in prs.slides)


def test_talk_deck_uses_glp_light_palette(tmp_path: Path):
    """讲稿按 Lumen-light：封面品牌绿渐变，正文页浅底，没有旧的深蓝底。"""
    out = render_talk(_deck(tmp_path), {}, [], tmp_path / "t.pptx")
    prs = Presentation(str(out))
    cover = prs.slides[0]._element.xml
    for c in ("00AA4F", "008E43", "006845"):
        assert c in cover, f"封面渐变缺 {c}"
    body = prs.slides[1]._element.xml
    assert "F9FAFC" in body                       # surface-page
    xml = _xml(prs)
    assert "0D162B" not in xml and "2563EB" not in xml, "还留着旧的深蓝主题"


def test_talk_bullets_marker_is_brand_green_and_text_is_gray(tmp_path: Path):
    out = render_talk(_deck(tmp_path), {}, [], tmp_path / "t.pptx")
    prs = Presentation(str(out))
    runs = [r for sh in prs.slides[1].shapes if sh.has_text_frame
            for p in sh.text_frame.paragraphs for r in p.runs]
    dots = [r for r in runs if r.text.strip() == "•"]
    assert dots and all(str(r.font.color.rgb) == "00AA4F" for r in dots)
    words = [r for r in runs if r.text.strip() not in ("", "•")]
    assert any(str(r.font.color.rgb) == "555B61" for r in words)


def test_talk_table_header_is_subtle_gray_without_vertical_lines(tmp_path: Path):
    out = render_talk(_deck(tmp_path), {}, [], tmp_path / "t.pptx")
    prs = Presentation(str(out))
    tables = [sh.table for sl in prs.slides for sh in sl.shapes if sh.has_table]
    assert tables
    tc = tables[0].cell(0, 0)._tc
    xml = tc.xml
    assert "F0F1F3" in xml                        # 表头 surface-subtle
    lnl = tc.tcPr.find(qn("a:lnL"))
    assert lnl is not None and lnl.find(qn("a:noFill")) is not None


def test_text_runs_set_chinese_font(tmp_path: Path):
    """不写 a:ea 的话中文会落到主题默认的宋体。"""
    out = render_talk(_deck(tmp_path), {}, [], tmp_path / "t.pptx")
    xml = _xml(Presentation(str(out)))
    assert 'typeface="Microsoft YaHei"' in xml
    assert 'typeface="Segoe UI"' in xml


def test_restore_deck_uses_brand_colors(tmp_path: Path):
    """按图还原也走品牌色：强调块品牌绿实底，普通块白卡片，页面浅底。"""
    plan = SlidePlan(title="展板", bands=[Band(columns=[Column(heading="栏", blocks=[
        Block(label="强调", accent="blue"), Block(label="普通")])])])
    out = render_deck([DeckItem(plan=plan)], tmp_path / "d.pptx")
    xml = Presentation(str(out)).slides[0]._element.xml
    assert "00AA4F" in xml and "FFFFFF" in xml and "F9FAFC" in xml


def test_render_talk_appendix_page_gets_notes(tmp_path: Path):
    """「未收录章节」那页的备注里放这些章节的原文 —— 讲的人需要时直接念。"""
    out = render_talk(_deck(tmp_path), {}, ["五、附录"], tmp_path / "t.pptx",
                      appendix_notes="【原文 · 五、附录】\n术语表")
    last = Presentation(str(out)).slides[-1]
    assert "术语表" in last.notes_slide.notes_text_frame.text
